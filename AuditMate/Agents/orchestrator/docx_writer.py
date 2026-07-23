"""DOCX writer for generating evidence reports."""
from pathlib import Path
from typing import Optional
from ..common import EvidenceSet, EVIDENCE_DIR


def write_docx(evidence_set: EvidenceSet) -> Optional[str]:
    """
    Write evidence set to DOCX file.
    
    Args:
        evidence_set: EvidenceSet to write
        
    Returns:
        Path to written DOCX or None if python-docx not available
    """
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        print("python-docx not installed. Skipping DOCX generation.")
        return None
    
    doc = Document()
    
    # Title
    title = doc.add_heading('Audit Evidence Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Metadata
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.style = 'Light Grid Accent 1'
    meta_table.cell(0, 0).text = "Request ID"
    meta_table.cell(0, 1).text = evidence_set.request_id
    meta_table.cell(1, 0).text = "Control Reference"
    meta_table.cell(1, 1).text = evidence_set.control_ref or "N/A"
    meta_table.cell(2, 0).text = "Instruction"
    meta_table.cell(2, 1).text = evidence_set.instruction[:100] + "..." if len(evidence_set.instruction) > 100 else evidence_set.instruction
    meta_table.cell(3, 0).text = "Items Collected"
    meta_table.cell(3, 1).text = str(len(evidence_set.items))
    
    # Draft notice
    draft = doc.add_paragraph()
    draft.add_run("DRAFT: ").bold = True
    draft.add_run("Evidence collection in progress. Subject to revision.")
    draft.style = 'Heading 3'
    
    # Audit points
    if evidence_set.audit_points:
        doc.add_heading('Audit Points', level=2)
        for point in evidence_set.audit_points:
            doc.add_paragraph(f"{point.id}: {point.text}", style='List Bullet')
    
    # Narrative
    if evidence_set.narrative:
        doc.add_heading('Narrative', level=2)
        doc.add_paragraph(evidence_set.narrative)
    
    # Evidence Items
    doc.add_heading('Evidence Items', level=2)
    
    for item in evidence_set.items:
        # Item heading
        doc.add_heading(f"{item.id}: {item.type.value}", level=3)
        
        # Summary
        doc.add_paragraph(f"Summary: {item.summary}")
        
        # Provenance
        prov = item.provenance
        doc.add_paragraph(f"Source: {prov.system} ({prov.method})", style='Normal')
        if prov.ref:
            doc.add_paragraph(f"Reference: {prov.ref}")
        
        # Content based on type
        if item.type.value == "code_flow" and item.flow:
            doc.add_heading('Code Flow', level=4)
            flow_table = doc.add_table(rows=1, cols=4)
            flow_table.style = 'Light Grid Accent 1'
            hdr_cells = flow_table.rows[0].cells
            hdr_cells[0].text = "Step"
            hdr_cells[1].text = "Artifact"
            hdr_cells[2].text = "Line"
            hdr_cells[3].text = "Detail"
            
            for step in item.flow:
                row_cells = flow_table.add_row().cells
                row_cells[0].text = step.type
                row_cells[1].text = step.ref
                row_cells[2].text = str(step.line or "")
                row_cells[3].text = step.detail
        
        elif item.type.value == "code_snippet" and item.content:
            doc.add_heading('Code', level=4)
            # Add code block (simple paragraph with monospace)
            code_para = doc.add_paragraph()
            code_run = code_para.add_run(item.content)
            code_run.font.name = 'Courier New'
            code_run.font.size = Pt(9)
        
        elif item.type.value == "screenshot" and item.image:
            doc.add_heading('Screenshot', level=4)
            try:
                # Try to embed image if file exists
                if Path(item.image).exists():
                    doc.add_picture(item.image, width=Inches(6))
            except:
                doc.add_paragraph(f"[Screenshot: {item.image}]")
        
        elif item.content:
            doc.add_paragraph(item.content)
        
        # Checks
        if item.checks:
            doc.add_heading('Checks', level=4)
            for check_name, check_result in item.checks.items():
                status = "✓" if check_result == "PASS" else "✗"
                doc.add_paragraph(f"{status} {check_name}: {check_result}", style='List Bullet')
        
        # Audit points related
        if item.audit_points:
            doc.add_paragraph(f"Related to: {', '.join(item.audit_points)}", style='Normal')
        
        doc.add_paragraph()  # Spacing
    
    # Relations
    if evidence_set.relations:
        doc.add_heading('Relations', level=2)
        rel_table = doc.add_table(rows=1, cols=3)
        rel_table.style = 'Light Grid Accent 1'
        hdr_cells = rel_table.rows[0].cells
        hdr_cells[0].text = "From"
        hdr_cells[1].text = "To"
        hdr_cells[2].text = "Kind"
        
        for rel in evidence_set.relations:
            row_cells = rel_table.add_row().cells
            row_cells[0].text = rel.from_id
            row_cells[1].text = rel.to_id
            row_cells[2].text = rel.kind
    
    # Save
    output_dir = Path(EVIDENCE_DIR)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{evidence_set.request_id}.docx"
    
    doc.save(str(output_path))
    return str(output_path)
